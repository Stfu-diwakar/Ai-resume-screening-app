"""
Utility to generate PDF and DOCX sample resumes for live upload testing.
"""
from pathlib import Path
import docx
import fitz  # PyMuPDF

SAMPLE_DIR = Path("data/sample_resumes")

def generate_docx_and_pdf():
    txt_files = list(SAMPLE_DIR.glob("*.txt"))
    for txt_file in txt_files:
        content = txt_file.read_text(encoding="utf-8")
        base_name = txt_file.stem

        # 1. Create DOCX
        doc = docx.Document()
        for line in content.splitlines():
            line_str = line.strip()
            if not line_str:
                doc.add_paragraph("")
            elif line_str.startswith("•"):
                doc.add_paragraph(line_str, style='List Bullet')
            elif any(line_str.startswith(h) for h in ["Professional Summary", "Technical Skills", "Work Experience", "Key Projects", "Education"]):
                doc.add_heading(line_str, level=2)
            elif line_str.startswith("Alex Morgan") or line_str.startswith("Sarah Lin") or line_str.startswith("Priya Sharma") or line_str.startswith("Jordan Reed") or line_str.startswith("Kevin Patel"):
                doc.add_heading(line_str, level=1)
            else:
                doc.add_paragraph(line_str)

        docx_path = SAMPLE_DIR / f"{base_name}.docx"
        doc.save(docx_path)
        print(f"Created DOCX: {docx_path}")

        # 2. Create PDF using PyMuPDF
        pdf_doc = fitz.open()
        pdf_page = pdf_doc.new_page(width=612, height=792)  # Standard Letter
        # Write text into rect
        rect = fitz.Rect(50, 50, 562, 742)
        pdf_page.insert_textbox(rect, content, fontsize=10, fontname="helv", color=(0.1, 0.1, 0.1))
        pdf_path = SAMPLE_DIR / f"{base_name}.pdf"
        pdf_doc.save(pdf_path)
        pdf_doc.close()
        print(f"Created PDF: {pdf_path}")

if __name__ == "__main__":
    generate_docx_and_pdf()
