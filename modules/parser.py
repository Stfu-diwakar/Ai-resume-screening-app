"""
Resume Parsing and Entity Extraction Module
Handles PDF, DOCX, and TXT parsing with contact extraction,
section segmentation, and spaCy / regex NLP preprocessing.
"""
import io
import re
from typing import Any, Dict, List, Optional
try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    docx = None
    DOCX_AVAILABLE = False

# Optional PyMuPDF import with fallback
try:
    import fitz  # PyMuPDF
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

# Optional spaCy import with fallback
try:
    import spacy
    try:
        nlp = spacy.load("en_core_web_sm")
    except Exception:
        # If model is not downloaded, create blank English pipeline
        nlp = spacy.blank("en")
except ImportError:
    nlp = None

from modules.taxonomy import extract_skills

# Regex patterns for contact information
EMAIL_REGEX = re.compile(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,7}\b')
PHONE_REGEX = re.compile(
    r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}\b'
)
LINKEDIN_REGEX = re.compile(r'(?:https?://)?(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+', re.IGNORECASE)
GITHUB_REGEX = re.compile(r'(?:https?://)?(?:www\.)?github\.com/[a-zA-Z0-9_-]+', re.IGNORECASE)
URL_REGEX = re.compile(r'https?://(?:[-\w.]|(?:%[\da-fA-F]{2}))+', re.IGNORECASE)

# Section header patterns
SECTION_PATTERNS = {
    "summary": re.compile(r'^\s*(professional\s+summary|executive\s+summary|summary|profile|about\s+me|objective)\s*$', re.IGNORECASE),
    "experience": re.compile(r'^\s*(work\s+experience|professional\s+experience|experience|employment\s+history|work\s+history|career\s+history)\s*$', re.IGNORECASE),
    "education": re.compile(r'^\s*(education|academic\s+background|qualifications|academic\s+history)\s*$', re.IGNORECASE),
    "projects": re.compile(r'^\s*(projects|key\s+projects|technical\s+projects|personal\s+projects|notable\s+projects)\s*$', re.IGNORECASE),
    "skills": re.compile(r'^\s*(technical\s+skills|skills\s+&\s+competencies|skills\s+and\s+tools|core\s+competencies|skills|technologies)\s*$', re.IGNORECASE),
    "certifications": re.compile(r'^\s*(certifications|licenses|certifications\s+and\s+awards|awards|achievements)\s*$', re.IGNORECASE)
}

def clean_text(text: str) -> str:
    """Normalize whitespace and remove non-printable characters."""
    if not text:
        return ""
    # Normalize unicode quotation marks and dashes
    text = text.replace('\u2019', "'").replace('\u2018', "'")
    text = text.replace('\u201c', '"').replace('\u201d', '"')
    text = text.replace('\u2013', '-').replace('\u2014', '-')
    text = text.replace('\u2022', '\n• ').replace('\u25cf', '\n• ')
    # Replace non-breaking spaces
    text = text.replace('\xa0', ' ')
    # Clean multiple consecutive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract plain text from PDF bytes using PyMuPDF (fitz)."""
    text = ""
    if PYMUPDF_AVAILABLE:
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page in doc:
                text += page.get_text("text") + "\n"
            doc.close()
            return clean_text(text)
        except Exception as e:
            print(f"PyMuPDF error: {e}")
    # Fallback to simple binary decode if fitz is unavailable or fails
    try:
        return clean_text(file_bytes.decode('utf-8', errors='ignore'))
    except Exception:
        return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx or zero-dependency zipfile/XML fallback."""
    if DOCX_AVAILABLE and docx is not None:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return clean_text("\n".join(full_text))
        except Exception as e:
            print(f"python-docx read error: {e}")

    # Built-in fallback: standard library zipfile + xml parser (zero external dependencies)
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            if 'word/document.xml' in zf.namelist():
                xml_content = zf.read('word/document.xml')
                tree = ET.fromstring(xml_content)
                texts = []
                for node in tree.iter():
                    if node.tag.endswith('}t') and node.text:
                        texts.append(node.text)
                if texts:
                    return clean_text(" ".join(texts))
    except Exception as e:
        print(f"zipfile xml fallback error: {e}")

    try:
        return clean_text(file_bytes.decode('utf-8', errors='ignore'))
    except Exception:
        return ""

def extract_contact_info(text: str) -> Dict[str, Any]:
    """Extract email, phone, linkedin, github, and URLs."""
    emails = EMAIL_REGEX.findall(text)
    phones = PHONE_REGEX.findall(text)
    # Filter plausible phone numbers (length >= 8 digits)
    valid_phones = []
    for p in phones:
        digits = re.sub(r'\D', '', p)
        if 8 <= len(digits) <= 15:
            valid_phones.append(p.strip())

    linkedin = LINKEDIN_REGEX.findall(text)
    github = GITHUB_REGEX.findall(text)
    all_urls = URL_REGEX.findall(text)

    return {
        "email": emails[0] if emails else None,
        "phone": valid_phones[0] if valid_phones else None,
        "linkedin": linkedin[0] if linkedin else None,
        "github": github[0] if github else None,
        "urls": list(set(all_urls))[:5]
    }

def extract_candidate_name(text: str, email: Optional[str] = None) -> str:
    """
    Extract the candidate's name from resume header.
    Heuristics:
    1. First clean non-empty line that looks like a person's name.
    2. spaCy PERSON entities in the top text.
    3. Infer from email.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return "Unknown Candidate"

    # Clean first line candidate
    first_line = re.split(r'(\bemail\b|\bphone\b|\||@)', lines[0], flags=re.IGNORECASE)[0].strip()
    if re.match(r'^[A-Z][a-zA-Z\'.\-]+(?:\s+[A-Z][a-zA-Z\'.\-]+){1,3}$', first_line):
        return first_line

    # spaCy NER on top of document if available
    top_text = "\n".join(lines[:4])
    if nlp is not None:
        try:
            doc = nlp(top_text)
            for ent in doc.ents:
                if ent.label_ == "PERSON":
                    clean_ent = re.split(r'(\bemail\b|\bphone\b|\n|\|)', ent.text, flags=re.IGNORECASE)[0].strip()
                    words = clean_ent.split()
                    if 2 <= len(words) <= 4:
                        if not re.search(r'(resume|curriculum|vitae|developer|engineer|manager|summary)', clean_ent, re.I):
                            return clean_ent
        except Exception:
            pass

    # If email is known, infer name from email prefix
    if email:
        prefix = email.split('@')[0]
        parts = re.split(r'[._-]', prefix)
        clean_parts = [p.capitalize() for p in parts if p.isalpha()]
        if len(clean_parts) >= 2:
            return " ".join(clean_parts[:2])

    return first_line[:30] if first_line else "Unknown Candidate"

def segment_sections(text: str) -> Dict[str, str]:
    """
    Split resume into structured sections based on recognized headings.
    """
    lines = text.splitlines()
    sections: Dict[str, List[str]] = {
        "header": [],
        "summary": [],
        "experience": [],
        "education": [],
        "projects": [],
        "skills": [],
        "certifications": [],
        "other": []
    }

    current_section = "header"

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        matched_new_section = False
        # Only test headings if short enough to be a section header
        if len(stripped) < 45 and not stripped.endswith('.'):
            for sec_name, pattern in SECTION_PATTERNS.items():
                if pattern.match(stripped):
                    current_section = sec_name
                    matched_new_section = True
                    break

        if not matched_new_section:
            sections[current_section].append(stripped)

    # Join lines back into text
    return {k: "\n".join(v).strip() for k, v in sections.items() if v}

def estimate_experience_years(text: str) -> float:
    """
    Extract mentioned years of experience from summary or experience text.
    Looks for phrases like '5+ years of experience', '3 years in software engineering', etc.
    """
    patterns = [
        r'(\d+(?:\.\d+)?)\+?\s*(?:years|yrs)\s*(?:of)?\s*(?:experience|exp)',
        r'(?:experience|exp)\s*(?:of)?\s*(\d+(?:\.\d+)?)\+?\s*(?:years|yrs)'
    ]
    years_found = []
    for pat in patterns:
        matches = re.findall(pat, text, re.IGNORECASE)
        for m in matches:
            try:
                y = float(m)
                if 0.5 <= y <= 35:
                    years_found.append(y)
            except ValueError:
                continue

    if years_found:
        return max(years_found)

    # Heuristic: count year ranges like 2019 - 2023
    year_ranges = re.findall(r'(20\d\d|19\d\d)\s*[-–—to]+\s*(20\d\d|present|current|now)', text, re.IGNORECASE)
    total_span = 0
    current_year = 2026
    for start, end in year_ranges:
        try:
            s = int(start)
            e = current_year if end.lower() in ['present', 'current', 'now'] else int(end)
            if 1980 <= s <= current_year and s <= e <= current_year:
                total_span += (e - s)
        except ValueError:
            continue

    if total_span > 0:
        return min(float(total_span), 25.0)

    return 1.0  # Default baseline

def parse_resume(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parse a resume from raw bytes (PDF, DOCX, or TXT).
    Returns complete parsed dictionary.
    """
    fname_lower = filename.lower()
    if fname_lower.endswith('.pdf'):
        raw_text = extract_text_from_pdf(file_bytes)
    elif fname_lower.endswith('.docx') or fname_lower.endswith('.doc'):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        # Default txt or string decode
        try:
            raw_text = clean_text(file_bytes.decode('utf-8', errors='ignore'))
        except Exception:
            raw_text = clean_text(str(file_bytes))

    contact = extract_contact_info(raw_text)
    name = extract_candidate_name(raw_text, contact.get("email"))
    sections = segment_sections(raw_text)
    skills_data = extract_skills(raw_text)
    exp_years = estimate_experience_years(raw_text)

    return {
        "filename": filename,
        "candidate_name": name,
        "email": contact.get("email"),
        "phone": contact.get("phone"),
        "linkedin": contact.get("linkedin"),
        "github": contact.get("github"),
        "urls": contact.get("urls", []),
        "sections": sections,
        "skills": skills_data["all_skills"],
        "skills_by_category": skills_data["by_category"],
        "skill_counts": skills_data["skill_counts"],
        "experience_years": exp_years,
        "full_text": raw_text,
        "char_count": len(raw_text),
        "word_count": len(raw_text.split())
    }

def parse_job_description(text: str, title: str = "Job Description") -> Dict[str, Any]:
    """
    Parse and analyze job description requirements.
    Extracts required skills, experience requirements, and sections.
    """
    cleaned = clean_text(text)
    skills_data = extract_skills(cleaned)
    exp_years = estimate_experience_years(cleaned)
    sections = segment_sections(cleaned)

    # Heuristic separation into required vs preferred skills
    required_keywords = ["must have", "required", "requirements", "qualifications", "minimum"]
    preferred_keywords = ["preferred", "nice to have", "plus", "bonus", "desirable", "good to have"]

    lines = cleaned.splitlines()
    req_skills: Set[str] = set()
    pref_skills: Set[str] = set()

    current_mode = "required"  # default
    for line in lines:
        lower_line = line.lower()
        if any(k in lower_line for k in preferred_keywords):
            current_mode = "preferred"
        elif any(k in lower_line for k in required_keywords):
            current_mode = "required"

        line_skills = extract_skills(line)["all_skills"]
        for s in line_skills:
            if current_mode == "preferred":
                pref_skills.add(s)
            else:
                req_skills.add(s)

    # All skills not in preferred are treated as required/core
    all_skills = skills_data["all_skills"]
    if not req_skills and not pref_skills:
        req_skills = set(all_skills)

    return {
        "title": title,
        "full_text": cleaned,
        "skills": all_skills,
        "required_skills": sorted(list(req_skills)),
        "preferred_skills": sorted(list(pref_skills - req_skills)),
        "skills_by_category": skills_data["by_category"],
        "min_experience_years": exp_years,
        "sections": sections
    }
